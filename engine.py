# engine.py - UNIFIED PACKING LIST ENGINE (FINAL: SHEET = OUTLET)
import os
import math
import copy
from collections import OrderedDict, defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side


class PackingListEngine:
    def __init__(self, input_file, output_file, progress_callback=None, master_wb=None):
        self.input_file = input_file
        self.output_file = output_file
        self.progress = progress_callback or (lambda msg: None)
        self.master_wb = master_wb
        self.wb = None
        self.MasterData = {}
        self.DOGroups = OrderedDict()
        self.TransformResult = OrderedDict()
        self.FinalResult = {}
        self.format_type = None
        self.target_sheet = ""
        self.logs = []

    def log(self, msg):
        self.logs.append(msg)
        self.progress(msg)

    def Txt(self, v):
        return "" if v is None else str(v).strip().upper()

    def TxtRaw(self, v):
        return "" if v is None else str(v).strip()

    def Val(self, v):
        try:
            if v is None or v == "": return 0
            return float(v)
        except:
            return 0

    def SafeSheetName(self, name):
        if not name: return "Sheet"
        bad = ['\\', '/', '*', '?', ':', '[', ']']
        for c in bad: name = name.replace(c, " ")
        return name[:31].strip()

    def NormalizeHeader(self, v):
        text = self.Txt(v)
        for old, new in [(" ", ""), (".", ""), ("_", ""), ("-", ""), ("/", ""), ("\\", "")]:
            text = text.replace(old, new)
        return text

    def NormalizeUnit(self, unit):
        unit_upper = self.Txt(unit)
        if unit_upper in ["DOZ", "DOZEN"]:
            return "DUS"
        return unit_upper

    def load_master_data(self):
        self.log("    Membaca Master Data...")
        master_sheet = next((s for s in self.wb.sheetnames if "Master Data" in s or "Master" in s), None)
        if not master_sheet:
            raise Exception(f"Sheet 'Master Data Packinglist' tidak ditemukan. Sheet yang ada: {self.wb.sheetnames}")
        
        ws = self.wb[master_sheet]
        for row in range(6, ws.max_row + 1):
            nama = self.Txt(ws[f"B{row}"].value)
            if not nama or nama in self.MasterData: continue
            self.MasterData[nama] = {
                "Nama": nama,
                "QtyMax": self.Val(ws[f"C{row}"].value),
                "Unit": self.Txt(ws[f"E{row}"].value),
                "Kategori": self.Txt(ws[f"F{row}"].value)
            }
        self.log(f"   ✅ Master Data: {len(self.MasterData)} item terbaca.")

    def detect_format(self):
        self.log("    Mendeteksi format Delivery Order...")
        sheets = self.wb.sheetnames
        rincian_sheet = next((s for s in sheets if "Rincian" in s or "Pemindahan" in s), None)
        do_sheet = next((s for s in sheets if "Delivery Order" in s or "DO Detail" in s), None)

        if rincian_sheet:
            self.format_type = "Rincian"
            self.target_sheet = rincian_sheet
        elif do_sheet:
            self.format_type = "DeliveryOrder"
            self.target_sheet = do_sheet
        else:
            raise Exception(f"Format DO tidak dikenali. Sheet: {sheets}")
        self.log(f"    Format: {self.format_type} (Sheet: '{self.target_sheet}')")

    def parse_rincian(self):
        self.log(f"    Parsing sheet: '{self.target_sheet}'...")
        ws = self.wb[self.target_sheet]
        header_keywords = {
            "DO": ["NOPEMINDAHAN", "NOPEMINDAHAN#", "NOMORPEMINDAHAN", "NOTRANSAKSI", "NOMORTRANSAKSI", "NO", "DO"],
            "TANGGAL": ["TANGGAL", "DATE", "TGL"],
            "GUDANG": ["GUDANGTUJUAN/DARI", "GUDANGTUJUAN", "GUDANGDARI", "GUDANG", "OUTLET"],
            "NAMA": ["NAMABARANG", "NAMAITEM", "ITEM", "BARANG"],
            "UNIT": ["SATUAN", "UNIT"],
            "QTY": ["KUANTITAS", "QTY", "QUANTITY", "JUMLAH"],
            "SHIPVIA": ["SHIPVIA", "SHIP", "VIA", "EKSPEDISI"]
        }
        header_row, header_map = None, {}
        for r in range(1, min(ws.max_row, 30) + 1):
            kandidat = {}
            for c in range(1, ws.max_column + 1):
                norm = self.NormalizeHeader(ws.cell(row=r, column=c).value)
                if not norm: continue
                for field, kws in header_keywords.items():
                    for kw in kws:
                        if norm == self.NormalizeHeader(kw):
                            kandidat[field] = c; break
            if "NAMA" in kandidat and "QTY" in kandidat:
                header_row, header_map = r, kandidat; break
                
        if not header_row:
            raise Exception(f"Header tidak ditemukan.")
        
        last_do, last_tgl, last_gudang, last_ship = None, None, None, None
        for r in range(header_row + 1, ws.max_row + 1):
            def cell(field):
                return ws.cell(row=r, column=header_map.get(field, 0)).value if field in header_map else None
            no_do = cell("DO")
            if no_do is not None:
                no_do = self.TxtRaw(no_do)
                if no_do: last_do = no_do
            else: no_do = last_do
            tgl = cell("TANGGAL")
            if tgl is not None: last_tgl = tgl
            else: tgl = last_tgl
            gudang = cell("GUDANG")
            if gudang is not None:
                g = self.TxtRaw(gudang)
                if g: last_gudang = g
            gudang = last_gudang
            ship = cell("SHIPVIA")
            if ship is not None:
                s = self.TxtRaw(ship)
                if s: last_ship = s
            ship = last_ship
            nama = cell("NAMA")
            if nama is None: continue
            nama = self.TxtRaw(nama)
            if not nama: continue
            qty = cell("QTY")
            try:
                if qty is None: continue
                qty = float(qty)
            except: continue
            if qty <= 0: continue
            unit = self.TxtRaw(cell("UNIT")) if "UNIT" in header_map else ""
            if not no_do: continue
            if no_do not in self.DOGroups:
                self.DOGroups[no_do] = {"Tanggal": tgl, "Outlet": gudang, "ShipVia": ship, "Items": []}
            self.DOGroups[no_do]["Items"].append({"Nama": nama, "Qty": qty, "Unit": unit})
        self.log(f"   ✅ DO terbaca: {len(self.DOGroups)}")

    def parse_delivery_order(self):
        self.log(f"    Parsing sheet: '{self.target_sheet}'...")
        ws = self.wb[self.target_sheet]
        header_rows = []
        for r in range(1, ws.max_row + 1):
            if self.Txt(ws.cell(row=r, column=3).value) == "NUMBER":
                header_rows.append(r)
        header_rows.append(ws.max_row + 1)
        if not header_rows or header_rows[0] > ws.max_row:
            raise Exception(f"Tidak ditemukan header 'NUMBER' di Kolom C.")
            
        for i in range(len(header_rows) - 1):
            start, end = header_rows[i], header_rows[i + 1]
            nomor_do = self.TxtRaw(ws.cell(row=start, column=6).value)
            outlet = self.TxtRaw(ws.cell(row=start, column=13).value)
            tanggal = ""
            if start + 1 < end:
                tanggal_val = ws.cell(row=start+1, column=6).value
                if tanggal_val: tanggal = tanggal_val
            ship_via = ""
            for col in range(14, 17):
                val = ws.cell(row=start, column=col).value
                if val:
                    ship_via = self.TxtRaw(val)
                    break
            if not nomor_do: continue
            self.DOGroups[nomor_do] = {"Outlet": outlet, "Tanggal": tanggal, "ShipVia": ship_via, "Items": []}
            for r in range(start, end):
                nama = self.TxtRaw(ws.cell(row=r, column=8).value)
                if "DUS" in self.Txt(nama) or not nama: continue
                if nama.upper() == "ITEM NAME": continue
                qty = self.Val(ws.cell(row=r, column=14).value)
                if qty <= 0: continue
                unit_raw = self.TxtRaw(ws.cell(row=r, column=16).value)
                unit_normalized = self.NormalizeUnit(unit_raw)
                self.DOGroups[nomor_do]["Items"].append({
                    "Nama": nama, "Qty": qty, "Unit": unit_normalized
                })
        self.log(f"   ✅ DO terbaca: {len(self.DOGroups)}")

    def get_rank(self, kategori):
        rank_map = {
            "DAGING": 1, "FROZEN VENDOR": 2, "FROZEN RECEH": 3,
            "KEJU RECEH": 4, "KEJU": 5, "KENTANG": 6, "SAUS": 7,
            "BUTTER": 8, "DRY RECEH": 9, "DRY VENDOR": 10,
            "MINYAK": 11, "GRILL BOX": 12, "HAMPERS": 13, "BUN": 999
        }
        return rank_map.get(kategori, 99)

    def transform(self):
        self.log("   ⚙️ Transforming data + Konversi UOM...")
        for do_no, data in self.DOGroups.items():
            hasil = []
            for item in data["Items"]:
                nama = self.TxtRaw(item["Nama"])
                key = self.Txt(nama)
                master = self.MasterData.get(key)
                if not master: continue
                if "DUS" in key: continue
                qty_max = master["QtyMax"]
                unit_master = master["Unit"]
                kategori_asli = master["Kategori"]
                unit_do = self.Txt(item["Unit"])
                qty_do = item["Qty"]
                if qty_max <= 0 or qty_do <= 0: continue
                nama_upper = key
                
                if unit_master == "DUS":
                    qty_real, qty_koli_max, unit_out = qty_do, qty_max, "DUS"
                elif unit_master == "PACK":
                    unit_out = "PACK"
                    if "DUS" in unit_do:
                        if "BEEF PATTY SMALL" in nama_upper or "BEEF PATTY LARGE" in nama_upper: isi_dus = 18
                        elif "THOUSAND ISLAND" in nama_upper: isi_dus = 20
                        else: isi_dus = qty_max
                        qty_real = qty_do * isi_dus
                    else: qty_real = qty_do
                    qty_koli_max = 20 if "THOUSAND ISLAND" in nama_upper else qty_max
                elif unit_master == "BKS":
                    unit_out = "BKS"
                    qty_real = qty_do * qty_max if "DUS" in unit_do else qty_do
                    qty_koli_max = qty_max
                else:
                    unit_out, qty_real, qty_koli_max = unit_master, qty_do, qty_max
                    
                if qty_koli_max <= 0: qty_koli_max = qty_max
                
                while qty_real > 0.0001:
                    qty_koli = min(qty_real, qty_koli_max)
                    status = "VENDOR" if abs(qty_koli - qty_koli_max) < 0.0001 else "RECEH"
                    if kategori_asli == "FROZEN": kat_out = "FROZEN VENDOR" if status == "VENDOR" else "FROZEN RECEH"
                    elif kategori_asli == "DRY": kat_out = "DRY VENDOR" if status == "VENDOR" else "DRY RECEH"
                    elif kategori_asli == "KEJU": kat_out = "KEJU" if status == "VENDOR" else "KEJU RECEH"
                    else: kat_out = kategori_asli
                        
                    hasil.append({
                        "NamaBarang": nama, "Qty": qty_koli, "Unit": unit_out, "Kategori": kat_out,
                        "Rank": self.get_rank(kat_out), "MaxKoli": qty_koli_max, "Status": status
                    })
                    qty_real -= qty_koli
                    
            hasil.sort(key=lambda x: (x["Rank"], x["NamaBarang"].upper()))
            self.TransformResult[do_no] = {
                "Outlet": data["Outlet"], "Tanggal": data.get("Tanggal", ""),
                "ShipVia": data.get("ShipVia", ""), "Items": hasil
            }
        self.log(f"   ✅ Transform: {len(self.TransformResult)} DO")

    def group_receh(self):
        self.log("   📦 Grouping RECEH...")
        MaxReceh = {"FROZEN RECEH": 35, "SAUS": 30, "KENTANG": 15, "PACKAGING": 60}
        for do_no, do_data in self.TransformResult.items():
            final, receh_groups = [], {}
            for item in do_data["Items"]:
                if item.get("Status") == "VENDOR": final.append(item)
                else: receh_groups.setdefault(item["Kategori"], []).append(item)
                    
            for kat, items in receh_groups.items():
                max_qty = MaxReceh.get(kat, 999)
                current_group, current_qty = [], 0
                for item in items:
                    qty = item["Qty"]
                    while qty > 0:
                        sisa = max_qty - current_qty
                        qty_masuk = min(qty, sisa)
                        new_item = item.copy(); new_item["Qty"] = qty_masuk
                        current_group.append(new_item)
                        current_qty += qty_masuk; qty -= qty_masuk
                        if current_qty >= max_qty:
                            final.extend(current_group); current_group, current_qty = [], 0
                if current_group: final.extend(current_group)
                    
            final.sort(key=lambda x: (x["Rank"], x["NamaBarang"].upper()))
            self.FinalResult[do_no] = {
                "Outlet": do_data["Outlet"], "Tanggal": do_data.get("Tanggal", ""), 
                "ShipVia": do_data.get("ShipVia", ""), "Items": final
            }
        self.log(f"   ✅ Final Result: {len(self.FinalResult)} DO")

    def export(self):
        self.log("   📝 Exporting Excel...")
        if self.master_wb is not None:
            out_wb = self.master_wb
        else:
            out_wb = Workbook()
            out_wb.remove(out_wb.active)

        if not self.FinalResult:
            self.log("   ⚠️ Tidak ada data DO yang berhasil diproses.")
            if self.master_wb is None:
                ws = out_wb.active; ws.title = "Info"
                ws["A1"] = "Tidak ada data Packing List yang bisa diproses."
                out_wb.save(self.output_file)
            return out_wb
            
        thin = Side(style="thin"); medium = Side(style="medium")
        border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
        border_medium = Border(left=medium, right=medium, top=medium, bottom=medium)
        
        for do_no, data in self.FinalResult.items():
            # ==========================================
            # INI BAGIAN PENTINGNYA: NAMA SHEET PAKAI OUTLET
            # ==========================================
            sheet_name = self.SafeSheetName(data["Outlet"])
            if not sheet_name:
                sheet_name = "DO"
                
            # Jika nama outlet sama (1 outlet terima 2 DO beda), tambah _2, _3 dst
            base_name = sheet_name; n = 2
            while sheet_name in out_wb.sheetnames:
                sheet_name = f"{base_name[:28]}_{n}"; n += 1
                
            ws = out_wb.create_sheet(title=sheet_name)
            ws.column_dimensions["A"].width = 10; ws.column_dimensions["B"].width = 38
            ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 25
            
            ws.merge_cells("A2:E2"); ws["A2"] = "PACKING LIST"
            ws["A2"].font = Font(bold=True, size=18); ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 30
            
            ws["A4"] = "DELIVERY\nORDER :"; ws["B4"] = do_no
            ws["D4"] = "Ship Via:"; ws["E4"] = data.get("ShipVia", "")
            ws["A5"] = "OUTLET :"; ws["B5"] = data["Outlet"]
            ws["A6"] = "DELIVERY\nDATE :"; ws["B6"] = data.get("Tanggal", "")
            
            for row in range(4, 7):
                ws.cell(row=row, column=1).font = Font(bold=True, size=12)
                ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                ws.cell(row=row, column=2).font = Font(size=12)
                ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center")
            if data.get("ShipVia"):
                ws.cell(row=4, column=4).font = Font(bold=True, size=12)
                ws.cell(row=4, column=5).font = Font(size=12)
            if data.get("Tanggal"):
                try: ws["B6"].number_format = "dd-mm-yyyy"
                except: pass
                
            ws.row_dimensions[4].height = 28; ws.row_dimensions[5].height = 20; ws.row_dimensions[6].height = 28
            
            hr = 9
            for col, h in enumerate(["NO KOLI", "NAMA BARANG", "QTY", "UNIT", "NOTES"], 1):
                c = ws.cell(row=hr, column=col, value=h)
                c.font = Font(bold=True, size=12); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = border_medium
            ws.row_dimensions[hr].height = 25
            
            r = 10
            for item in data["Items"]:
                ws.cell(row=r, column=2, value=item["NamaBarang"]).font = Font(size=12)
                ws.cell(row=r, column=3, value=item["Qty"]).font = Font(size=12)
                ws.cell(row=r, column=4, value=item["Unit"]).font = Font(size=12)
                for col in range(1, 6):
                    cell = ws.cell(row=r, column=col); cell.border = border_thin; cell.font = Font(size=12)
                ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=4).alignment = Alignment(horizontal="center")
                r += 1
                
            ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            
        if self.master_wb is None:
            out_wb.save(self.output_file)
        self.log(f"   ✅ Sheet berhasil ditambahkan.")
        return out_wb

    def run(self):
        self.logs = []
        self.wb = load_workbook(self.input_file, data_only=False)
        self.load_master_data()
        self.detect_format()
        if self.format_type == "Rincian": self.parse_rincian()
        else: self.parse_delivery_order()
        self.transform()
        self.group_receh()
